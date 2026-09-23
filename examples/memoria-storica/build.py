"""Genera la memoria storica finta dell'esempio «Memoria storica».

Tutto è inventato: aziende, persone, quote di fatturato, clausole.
Uso: ../../env/bin/python build.py   (serve Chrome per il PDF)
"""

import sys
from pathlib import Path

import openpyxl
from docx import Document

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent / "_shared"))
from fixtures import eml, pdf  # noqa: E402

# Il CEO chiede una bozza di contratto per Nordika Mobility. Quello che nessun
# modello di frontiera può sapere è scritto solo qui: Nordika è partner diretto di
# Veloce Automotive, cliente storico, e l'NDA con Veloce impone di avvisarli prima.
MEM = HERE / "Pratiche" / "Memoria-Storica"
COM = HERE / "Pratiche" / "Commerciale"


def memoria() -> None:
    MEM.mkdir(parents=True, exist_ok=True)
    COM.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clienti storici"
    ws.append(["Cliente", "Dal", "Quota fatturato 2025", "Referente", "Note"])
    ws.append(["Veloce Automotive S.p.A.", 2014, "38%", "ing. Davide Ferrario",
               "Banchi di test BMS per la piattaforma elettrica Kappa. NDA VA-2019-07."])
    ws.append(["Officine Brembo Nord S.r.l.", 2017, "11%", "Sara Longhi", "Collaudo centraline freno."])
    ws.append(["Mobilità Adriatica S.p.A.", 2021, "7%", "Luca Pini", "Retrofit linee di collaudo."])
    wb.save(MEM / "Anagrafica_clienti_storici.xlsx")

    d = Document()
    d.add_heading("Verbale riunione con Veloce Automotive — 12 marzo 2026", 1)
    d.add_paragraph("Presenti: CEO Meccatronica Lariana; ing. Davide Ferrario e Giorgio Sala per Veloce Automotive.")
    d.add_paragraph(
        "Veloce comunica, in via riservata, che la piattaforma elettrica Kappa sarà sviluppata insieme a "
        "Nordika Mobility GmbH, partner commerciale diretto con accordo di fornitura in esclusiva dal 2025. "
        "Nordika riceverà le specifiche dei banchi di test BMS che realizziamo per Veloce."
    )
    d.add_paragraph(
        "Veloce chiede che qualunque richiesta di Nordika o di altri partner della piattaforma Kappa venga "
        "segnalata prima di rispondere, come previsto dalla clausola 7.3 dell'NDA VA-2019-07. "
        "Ferrario: «Se scopriamo che lavorate per i nostri partner senza dircelo, la fiducia finisce lì.»"
    )
    d.save(MEM / "Verbale_Veloce_2026-03-12.docx")

    pdf(
        """<h1>Accordo di riservatezza VA-2019-07</h1>
<p class="small">tra Veloce Automotive S.p.A. e Meccatronica Lariana S.r.l. · rinnovato il 15 gennaio 2025</p>
<p>1. Le Informazioni Riservate comprendono le specifiche dei banchi di test BMS, la roadmap della
piattaforma Kappa e l'identità dei partner coinvolti.</p>
<p><b>7.3 Partner della piattaforma.</b> Prima di accettare incarichi da partner commerciali di Veloce
Automotive coinvolti nella piattaforma Kappa, Meccatronica Lariana ne dà comunicazione scritta a Veloce,
che può opporsi entro 15 giorni.</p>
<p>9. La violazione dell'articolo 7 è giusta causa di recesso dal contratto quadro.</p>""",
        MEM / "NDA_Veloce_VA-2019-07.pdf",
    )

    (MEM / "Note_direzione.md").write_text(
        "# Note della direzione (non condividere)\n\n"
        "- Veloce Automotive vale il 38% del fatturato 2025: è il cliente da proteggere prima di tutti.\n"
        "- Rinnovo del contratto quadro con Veloce a dicembre 2026.\n"
        "- Con Nordika Mobility c'è stato un solo contatto, a una fiera nel 2024.\n",
        encoding="utf-8",
    )

    eml(
        COM / "CEO_bozza_contratto_Nordika.eml",
        "CEO <ceo@meccatronica-lariana.example>",
        "Ufficio Vendite <vendite@meccatronica-lariana.example>",
        "Bozza contratto Nordika Mobility",
        "Tue, 22 Sep 2026 18:05:00 +0200",
        "Ciao a tutti,\n\nNordika Mobility GmbH ci chiede 4 banchi di test BMS, consegna primo trimestre "
        "2027. Preparate una bozza di contratto entro venerdì, condizioni standard.\n\nGrazie\nIl CEO\n",
    )


if __name__ == "__main__":
    memoria()
    print("ok:", *sorted(p.relative_to(HERE) for p in (HERE / "Pratiche").rglob("*.*") if ".annona-cache" not in str(p)), sep="\n  ")
